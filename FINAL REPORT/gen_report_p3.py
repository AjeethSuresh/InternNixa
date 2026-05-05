import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(__file__)
DIAG = os.path.join(BASE, "diagrams")
OUT  = os.path.join(BASE, "INTERNIXA_FINAL_REPORT.docx")

doc = Document()

# ─── Page setup A4 ───────────────────────────────────────────────────────────
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.54); sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.17); sec.right_margin = Cm(2.54)

# ─── Style helpers ───────────────────────────────────────────────────────────
def para(text="", bold=False, italic=False, size=12, align="left", color=None, space_before=0, space_after=6, font="Times New Roman"):
    p = doc.add_paragraph()
    p.alignment = {"left":WD_ALIGN_PARAGRAPH.LEFT,"center":WD_ALIGN_PARAGRAPH.CENTER,
                   "justify":WD_ALIGN_PARAGRAPH.JUSTIFY,"right":WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = Pt(18)
    if text:
        run = p.add_run(text)
        run.bold = bold; run.italic = italic
        run.font.name = font; run.font.size = Pt(size)
        if color: run.font.color.rgb = RGBColor(*color)
    return p

def chapter_heading(num, title):
    doc.add_page_break()
    para(f"CHAPTER {num}", bold=True, size=14, align="center", space_before=24, space_after=4)
    para(title.upper(), bold=True, size=14, align="center", space_before=0, space_after=18)

def section(num, title):
    para(f"{num}  {title}", bold=True, size=12, align="left", space_before=12, space_after=6)

def subsection(num, title):
    para(f"{num}  {title}", bold=True, italic=True, size=12, align="left", space_before=8, space_after=4)

def body(text):
    para(text, size=12, align="justify", space_before=0, space_after=6)

def fig(path, caption):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(5.5))
    para(caption, italic=True, size=11, align="center", space_before=2, space_after=10)

def add_table(headers, rows, caption):
    para(caption, bold=True, italic=True, size=11, align="center", space_before=8, space_after=4)
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True; run.font.size = Pt(10)
        tc = hdr[i]._tc; tcp = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'),'BFD7FF'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:val'),'clear')
        tcp.append(shd)
    for ri,row in enumerate(rows):
        cells = t.rows[ri+1].cells
        for ci,cell in enumerate(row):
            cells[ci].text = str(cell)
            for run in cells[ci].paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════════
#  TITLE PAGE
# ═══════════════════════════════════════════════════════════════════════════
para("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", bold=True, size=14, align="center", space_before=60, space_after=8)
para("Final Year Project Report", bold=False, size=12, align="center", space_before=4, space_after=40)
para("INTERNIXA", bold=True, size=20, align="center", space_before=10, space_after=8, color=(79,70,229))
para("AN AI-MONITORED E-LEARNING ECOSYSTEM WITH REAL-TIME\nENGAGEMENT INTELLIGENCE AND GENERATIVE AI ASSISTANCE",
     bold=True, size=14, align="center", space_before=4, space_after=40)
para("Submitted by", size=12, align="center", space_before=30, space_after=8)
for name in ["[STUDENT NAME 1]","[STUDENT NAME 2]","[STUDENT NAME 3]","[STUDENT NAME 4]"]:
    para(name, bold=True, size=12, align="center", space_before=2, space_after=2)
para("Academic Year: 2025–2026", size=12, align="center", space_before=20, space_after=4)
para("Degree: Bachelor of Engineering – Computer Science and Engineering", size=12, align="center")
doc.add_page_break()

# CERTIFICATE
para("BONAFIDE CERTIFICATE", bold=True, size=14, align="center", space_before=40, space_after=20)
body("This is to certify that the project entitled INTERNIXA: AN AI-MONITORED E-LEARNING ECOSYSTEM WITH REAL-TIME ENGAGEMENT INTELLIGENCE AND GENERATIVE AI ASSISTANCE is a bonafide record of work done by the above-mentioned students of the Department of Computer Science and Engineering in partial fulfilment of the requirements for the award of the degree of Bachelor of Engineering in Computer Science and Engineering during the academic year 2025–2026.")
para("\n\nGuide: [GUIDE NAME]                                HOD: [HOD NAME]", size=12, align="left", space_before=40)
para("Designation: Assistant Professor                   Head of Department", size=12, align="left")
doc.add_page_break()

# DECLARATION
para("DECLARATION", bold=True, size=14, align="center", space_before=40, space_after=20)
body("We hereby declare that the project work entitled INTERNIXA: AN AI-MONITORED E-LEARNING ECOSYSTEM WITH REAL-TIME ENGAGEMENT INTELLIGENCE AND GENERATIVE AI ASSISTANCE submitted to the Department of Computer Science and Engineering is a record of original work done by us under the guidance of [GUIDE NAME]. This project work has not been submitted previously for the award of any degree or diploma.")
para("\n\nSignatures:", size=12, space_before=30)
for name in ["[STUDENT NAME 1]","[STUDENT NAME 2]","[STUDENT NAME 3]","[STUDENT NAME 4]"]:
    para(name, size=12, space_before=8)
doc.add_page_break()

# ACKNOWLEDGEMENT
para("ACKNOWLEDGEMENT", bold=True, size=14, align="center", space_before=40, space_after=20)
body("We express our sincere gratitude to [GUIDE NAME], our project guide, for their invaluable guidance, constant encouragement, and insightful suggestions throughout the development of this project. Their technical expertise and academic mentorship have been instrumental in shaping the direction and quality of this work.")
body("We are deeply thankful to [HOD NAME], Head of the Department of Computer Science and Engineering, for providing us with the necessary resources, infrastructure, and a stimulating academic environment to carry out this project.")
body("We also extend our heartfelt appreciation to the management and faculty of our institution for their continued support, and to our families and friends whose unwavering encouragement gave us the strength to persevere through the challenges of this endeavour.")
doc.add_page_break()

# ABSTRACT
para("ABSTRACT", bold=True, size=14, align="center", space_before=40, space_after=20)
body("The proliferation of online learning platforms has revolutionized access to education; however, it has simultaneously introduced a critical accountability gap. The \"Background Video\" problem — wherein students play educational content without actively engaging with it — renders conventional e-learning platforms fundamentally ineffective. Internixa is a full-stack, AI-augmented e-learning ecosystem engineered specifically to resolve this challenge by enforcing genuine, verified engagement at every stage of the learning process.")
body("The platform's core innovation lies in its browser-native Artificial Intelligence monitoring layer, implemented using Google's MediaPipe Face Mesh library. This module processes the student's live camera feed at 30 frames per second, extracting 468 three-dimensional facial landmarks. The Eye Aspect Ratio (EAR) formula — computed from key eyelid landmark distances — detects blink events and prolonged eye closure. Simultaneously, a gaze deviation algorithm calculates the horizontal displacement of the nose tip relative to the eye midpoint to determine if the student's attention has drifted away from the screen. When distraction is confirmed for three or more consecutive frames, the course video pauses automatically; it resumes only when active attention is restored.")
body("Internixa also integrates a Retrieval-Augmented Generation (RAG) pipeline to deliver a context-locked AI chatbot. Course video transcripts are chunked, embedded using Google's text-embedding-004 model, and stored in a Pinecone vector database. When a student submits a query, the system retrieves the most semantically similar transcript segments and passes them as context to the Google Gemini 1.5 Flash large language model, ensuring that all AI responses are grounded exclusively in course content.")
body("Additional features include live WebRTC-based video meeting rooms with a host engagement HUD powered by Socket.IO, a gamified Focus Points system with a competitive leaderboard, auto-generated PDF certificates issued only upon achieving a verified engagement score of seventy percent or higher, AI-generated Smart Study Sheets, and a dedicated Recruiter Portal for internship management. The platform is deployed on Vercel (frontend) and Railway (backend), with MongoDB serving as the primary database. Internixa demonstrates that AI monitoring, when combined with thoughtful UX design, can transform passive digital consumption into disciplined, accountable, and measurably effective learning.")
doc.add_page_break()

# TABLE OF CONTENTS
para("TABLE OF CONTENTS", bold=True, size=14, align="center", space_before=20, space_after=16)
toc_items = [
    ("Abstract","iv"),("List of Figures","vi"),("List of Tables","vii"),("List of Abbreviations","viii"),
    ("CHAPTER 1 – INTRODUCTION","1"),
    ("  1.1  Background","2"),("  1.2  Problem Statement","4"),("  1.3  Objectives","6"),
    ("CHAPTER 2 – LITERATURE REVIEW","8"),
    ("  2.1  Introduction","8"),("  2.2  Review of Existing Work","9"),("  2.3  Comparative Analysis","14"),
    ("CHAPTER 3 – PROPOSED SYSTEM","16"),
    ("  3.1  Existing System","16"),("  3.2  Proposed System","18"),
    ("  3.3  System Architecture","20"),("  3.4  Methodology / Algorithm","24"),
    ("  3.5  Module Description","30"),
    ("CHAPTER 4 – IMPLEMENTATION & RESULTS","36"),
    ("  4.1  Tools & Technologies Used","36"),("  4.2  System Implementation","38"),
    ("  4.3  Working Model","40"),("  4.4  Screenshots / UI Walkthrough","42"),
    ("  4.5  Results & Outputs","48"),("  4.6  Performance Analysis","50"),
    ("CHAPTER 5 – CONCLUSION","52"),
    ("  5.1  Final Outcome","52"),("  5.2  Future Enhancements","54"),
    ("Appendices","57"),("References","61"),
]
for title, pg in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.5))
    run = p.add_run(f"{title}\t{pg}")
    run.font.name = "Times New Roman"; run.font.size = Pt(11)
    if "CHAPTER" in title: run.bold = True
doc.add_page_break()

# LIST OF FIGURES
para("LIST OF FIGURES", bold=True, size=14, align="center", space_before=20, space_after=12)
figs_list = [
    ("Fig 1.1","Internixa System Overview","2"),
    ("Fig 3.1","Traditional LMS Architecture","16"),
    ("Fig 3.2","Internixa 5-Layer Architecture","20"),
    ("Fig 3.3","DFD Level 0 – Context Diagram","22"),
    ("Fig 3.4","DFD Level 1 – Detailed Flow","23"),
    ("Fig 3.5","MediaPipe Face Mesh (468 Landmarks)","25"),
    ("Fig 3.6","Eye Aspect Ratio (EAR) Landmark Diagram","26"),
    ("Fig 3.7","Gaze Detection Algorithm Flow","27"),
    ("Fig 3.8","RAG Pipeline Flow","28"),
    ("Fig 3.9","WebRTC Signaling Diagram","29"),
    ("Fig 3.10","Certificate Generation Flow","35"),
    ("Fig 4.1","System Requirements Overview","36"),
    ("Fig 4.2","Engagement Score Calculation Flow","48"),
    ("Fig 4.3","Focus Points Algorithm Flowchart","49"),
]
add_table(["Figure No.","Title","Page No."],figs_list,"")
doc.add_page_break()

# LIST OF TABLES
para("LIST OF TABLES", bold=True, size=14, align="center", space_before=20, space_after=12)
add_table(["Table No.","Title","Page No."],
    [("Table 2.1","Comparative Analysis of E-Learning Platforms","14"),
     ("Table 4.1","Tools and Technologies Used","36"),
     ("Table 4.2","Performance Benchmarks","50")],"")
doc.add_page_break()

# ABBREVIATIONS
para("LIST OF ABBREVIATIONS", bold=True, size=14, align="center", space_before=20, space_after=12)
abbrevs = [
    ("AI","Artificial Intelligence"),("EAR","Eye Aspect Ratio"),
    ("RAG","Retrieval-Augmented Generation"),("LLM","Large Language Model"),
    ("JWT","JSON Web Token"),("FP","Focus Points"),
    ("WebRTC","Web Real-Time Communication"),("ASGI","Asynchronous Server Gateway Interface"),
    ("API","Application Programming Interface"),("DFD","Data Flow Diagram"),
    ("FPS","Frames Per Second"),("HUD","Heads-Up Display"),
    ("UI","User Interface"),("UX","User Experience"),
    ("PDF","Portable Document Format"),("LMS","Learning Management System"),
    ("MOOC","Massive Open Online Course"),("NoSQL","Not Only SQL"),
    ("STUN","Session Traversal Utilities for NAT"),("SDP","Session Description Protocol"),
]
add_table(["Abbreviation","Full Form"],abbrevs,"")

print("Front matter done.")
doc.save(OUT)
print(f"Saved intermediate to {OUT}")
